"""Unit tests for AMFI monthly flow report ingestion and parsing.

Verifies:
- Format detection (OOXML, BIFF8, invalid)
- Modern layout parsing (OOXML and BIFF8)
- Legacy multi-sheet layout parsing (2018)
- Malformed layout error diagnostics
- Units confirmation and safety
- Non-leakage available_from verification
- Immutability of raw downloaded files
- Zero inference of gross redemptions from net flows
- Strict contract validation
"""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
import pytest
import openpyxl
import pandas as pd

from fundflow.contracts import validate_table, COLUMNS
from fundflow.ingest.amfi_flows import (
    detect_file_format,
    parse_report,
    download_report,
    _clean_numeric,
    _parse_month_string,
    _parse_release_date_to_iso,
)

ROOT = Path(__file__).resolve().parents[1]
RAW_AMFI_DIR = ROOT / "data" / "raw" / "amfi"


def test_clean_numeric_helper():
    assert _clean_numeric("1,234.56") == 1234.56
    assert _clean_numeric(45.6) == 45.6
    assert _clean_numeric("-") is None
    assert _clean_numeric("N.A.") is None
    assert _clean_numeric(None) is None
    assert _clean_numeric("invalid") is None


def test_month_parsing_helper():
    assert _parse_month_string("Monthly Report for September 2024") == "2024-09"
    assert _parse_month_string("Consolidated Data for April 2019") == "2019-04"
    assert _parse_month_string("amsep2026repo.xls") == "2026-09"
    assert _parse_month_string("amjan2021repo.xls") == "2021-01"


def test_release_date_parsing_helper():
    assert _parse_release_date_to_iso("Released on 10-Oct-2024") == "2024-10-10T00:00:00Z"
    assert _parse_release_date_to_iso("Released on 08.10.2018") == "2018-10-08T00:00:00Z"
    assert _parse_release_date_to_iso("No date here") is None


def test_detect_file_format_ooxml(tmp_path):
    f = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    wb.save(f)
    assert detect_file_format(f) == "OOXML"


def test_detect_file_format_biff8():
    biff8_path = RAW_AMFI_DIR / "amsep2026repo.xls"
    if biff8_path.is_file():
        assert detect_file_format(biff8_path) == "BIFF8"


def test_detect_file_format_empty_fails(tmp_path):
    empty_f = tmp_path / "empty.xls"
    empty_f.write_bytes(b"")
    with pytest.raises(ValueError, match="empty"):
        detect_file_format(empty_f)


def test_detect_file_format_unknown_magic_fails(tmp_path):
    corrupt_f = tmp_path / "corrupt.xls"
    corrupt_f.write_bytes(b"NOT_A_VALID_EXCEL_HEADER_BYTES")
    with pytest.raises(ValueError, match="Unrecognized file format"):
        detect_file_format(corrupt_f)


def test_parse_modern_ooxml_report():
    rep_path = RAW_AMFI_DIR / "amsep2024repo.xls"
    if not rep_path.is_file():
        pytest.skip("amsep2024repo.xls not present in data/raw/amfi")

    # Hash before parsing
    hash_before = hashlib.sha256(rep_path.read_bytes()).hexdigest()

    df = parse_report(rep_path)

    # Invariant: untouched download
    hash_after = hashlib.sha256(rep_path.read_bytes()).hexdigest()
    assert hash_before == hash_after, "Parser mutated the raw report file!"

    # Verify contract
    assert tuple(df.columns) == COLUMNS["flows_category_month"]
    validate_table("flows_category_month", df)

    assert len(df) >= 6
    assert (df["month"] == "2024-09").all()
    assert "Large Cap" in df["category"].values
    assert "Mid Cap" in df["category"].values
    assert "Small Cap" in df["category"].values
    assert "Flexi Cap" in df["category"].values

    # Invariant: available_from is evidence-backed, not month-end
    assert not any(df["available_from"].str.startswith("2024-09-30"))
    assert (df["available_from"] == "2024-10-10T00:00:00Z").all()

    # Invariant: arithmetic identity holds (net = gross_inflow - gross_redemption)
    diff = (df["gross_inflow_cr"] - df["gross_redemption_cr"]) - df["net_flow_cr"]
    assert (diff.abs() < 0.05).all()


def test_parse_modern_biff8_report():
    rep_path = RAW_AMFI_DIR / "amsep2026repo.xls"
    if not rep_path.is_file():
        pytest.skip("amsep2026repo.xls not present in data/raw/amfi")

    hash_before = hashlib.sha256(rep_path.read_bytes()).hexdigest()

    df = parse_report(rep_path)

    hash_after = hashlib.sha256(rep_path.read_bytes()).hexdigest()
    assert hash_before == hash_after

    assert tuple(df.columns) == COLUMNS["flows_category_month"]
    validate_table("flows_category_month", df)

    assert len(df) >= 6
    assert (df["month"] == "2026-09").all()
    assert "Large Cap" in df["category"].values
    assert "Mid Cap" in df["category"].values
    assert "Aggressive Hybrid" in df["category"].values

    assert not any(df["available_from"].str.startswith("2026-09-30"))


def test_parse_legacy_biff8_report():
    rep_path = RAW_AMFI_DIR / "amsep2018repo.xls"
    if not rep_path.is_file():
        pytest.skip("amsep2018repo.xls not present in data/raw/amfi")

    df = parse_report(rep_path)
    assert tuple(df.columns) == COLUMNS["flows_category_month"]
    validate_table("flows_category_month", df)
    assert (df["month"] == "2018-09").all()
    assert "Liquid" in df["category"].values


def test_malformed_layout_missing_headers_fails(tmp_path):
    bad_f = tmp_path / "bad_headers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sep 24"
    ws.append(["Title", "No Inflow Or Redemption Header Here"])
    ws.append(["Random Data", 1234])
    wb.save(bad_f)

    with pytest.raises(ValueError, match="header keywords .* missing"):
        parse_report(bad_f, available_from="2024-10-10T00:00:00Z")


def test_malformed_layout_unconfirmed_units_fails(tmp_path):
    bad_units_f = tmp_path / "bad_units.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sep 24"
    ws.append(["Report for the month of September 2024 (in Thousands)"])
    ws.append(["Scheme Name", "Funds Mobilized", "Repurchase/ Redemption", "Net Inflow", "Net Assets Under Management"])
    ws.append(["Large Cap Fund", 100, 50, 50, 1000])
    wb.save(bad_units_f)

    with pytest.raises(ValueError, match="Unconfirmed units"):
        parse_report(bad_units_f, available_from="2024-10-10T00:00:00Z")


def test_missing_available_from_fails(tmp_path):
    no_avail_f = tmp_path / "no_avail.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sep 24"
    ws.append(["Report for the month of September 2024 (INR in crore)"])
    ws.append(["Scheme Name", "Funds Mobilized", "Repurchase/ Redemption", "Net Inflow", "Net Assets Under Management"])
    ws.append(["Large Cap Fund", 100, 50, 50, 1000])
    wb.save(no_avail_f)

    with pytest.raises(ValueError, match="Missing required publication lag evidence"):
        parse_report(no_avail_f, available_from="")


def test_no_gross_redemption_inferred_from_net_flows(tmp_path):
    missing_red_f = tmp_path / "missing_red.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sep 24"
    ws.append(["Report for the month of September 2024 (INR in crore)"])
    ws.append(["Released on 10-Oct-2024"])
    ws.append(["Scheme Name", "Funds Mobilized", "Repurchase/ Redemption", "Net Inflow", "Net Assets Under Management"])
    ws.append(["Large Cap Fund", 100, "-", 50, 1000])
    wb.save(missing_red_f)

    with pytest.raises(ValueError, match="Missing observed gross flow values"):
        parse_report(missing_red_f)


def test_download_report_invalid_url():
    with pytest.raises(ValueError, match="credentials"):
        download_report("https://user:pass@example.com/report.xls")
    with pytest.raises(ValueError, match="query"):
        download_report("https://example.com/report.xls?key=secret")


def test_clean_numeric_distinguishes_zero_from_missing():
    assert _clean_numeric(0) == 0.0
    assert _clean_numeric(0.0) == 0.0
    assert _clean_numeric("0") == 0.0
    assert _clean_numeric("0.0") == 0.0
    assert _clean_numeric("-") is None
    assert _clean_numeric("N.A.") is None
    assert _clean_numeric("") is None
    assert _clean_numeric(None) is None


def test_parse_flexi_cap_introduction_in_jan_2021():
    rep_path = RAW_AMFI_DIR / "amjan2021repo.xls"
    if not rep_path.is_file():
        pytest.skip("amjan2021repo.xls not present")
    df = parse_report(rep_path)
    assert (df["month"] == "2021-01").all()
    assert "Flexi Cap" in df["category"].values
    flexi_row = df[df["category"] == "Flexi Cap"].iloc[0]
    assert flexi_row["gross_inflow_cr"] > 0
    assert flexi_row["gross_redemption_cr"] > 0
    assert pytest.approx(flexi_row["gross_inflow_cr"] - flexi_row["gross_redemption_cr"]) == flexi_row["net_flow_cr"]


def test_all_eight_historical_reports_reconciliation():
    files = [
        "amapr2018repo.xls", "amsep2018repo.xls", "amapr2019repo.xls", "amsep2020repo.xls",
        "amjan2021repo.xls", "amsep2022repo.xls", "amsep2024repo.xls", "amsep2026repo.xls"
    ]
    for filename in files:
        rep_p = RAW_AMFI_DIR / filename
        if not rep_p.is_file():
            continue
        df = parse_report(rep_p)
        validate_table("flows_category_month", df)
        diff = (df["gross_inflow_cr"] - df["gross_redemption_cr"]) - df["net_flow_cr"]
        assert (diff.abs() < 0.01).all(), f"Accounting discrepancy in {filename}"
